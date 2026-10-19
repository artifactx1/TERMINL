#!/usr/bin/env python3
"""Build a private, auditable GTD/FCFS manifest from TERMINL collab sheets.

Requires openpyxl. Source XLSX files and sources.json live in the private source
directory; neither the wallet manifest nor audit is suitable for Git.
"""

import argparse
import csv
import json
import re
from collections import Counter
from pathlib import Path

import openpyxl

WALLET = re.compile(r"(?i)(?<![0-9a-z])0x[0-9a-f]{40}(?![0-9a-z])")
TIER = re.compile(r"(?i)\b(GTD|FCFS|GUARANTEED)\b")


def tier(value):
    match = TIER.search(str(value or ""))
    if match:
        return "gtd" if match.group(1).lower() in ("gtd", "guaranteed") else "fcfs"
    return "gtd" if str(value or "").strip().lower() == "team" else None


def wallets(value):
    return [match.group().lower() for match in WALLET.finditer(str(value or ""))]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source_dir", type=Path)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output_dir", type=Path)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    sources = json.loads((args.source_dir / "sources.json").read_text())
    evidence = []
    issues = []
    workbook = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True).active
    workbook_tiers = Counter()
    for row_number, row in enumerate(workbook.iter_rows(values_only=True), 1):
        address = wallets(row[1]) if len(row) > 1 else []
        if not address:
            if len(row) > 1 and "0x" in str(row[1]).lower():
                issues.append({"source": "local workbook", "row": row_number, "value": str(row[1])[:100], "problem": "invalid wallet"})
            continue
        kind = str(row[2] or "").strip().lower() if len(row) > 2 else ""
        if kind in ("gtd", "fcfs"):
            for value in address:
                evidence.append({"source_id": "local workbook", "community": str(row[3] or "").strip(),
                                 "tab": workbook.title, "row": row_number, "column": 2,
                                 "address": value, "kind": kind, "classification": "workbook explicit"})
                workbook_tiers[kind] += 1
        else:
            issues.append({"source": "local workbook", "row": row_number, "value": str(row[1])[:100],
                           "problem": "unconfirmed tier: " + kind})

    for source in sources:
        path = args.source_dir / (source["id"] + ".xlsx")
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet in book:
            tab_kind = tier(sheet.title)
            section = "gtd" if source["id"] == "1PAbwbdkHpuljTl6kwlBVuRtyfF8Q-zcw274o8Mji718" else None
            column_tiers = {}
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                found = [(column, value) for column, cell in enumerate(row, 1)
                         for value in wallets(cell)]
                labels = [(column, tier(cell)) for column, cell in enumerate(row, 1) if tier(cell)]
                if not found and len(labels) > 1:
                    column_tiers = dict(labels)
                elif labels:
                    section = labels[-1][1]
                for column, address in found:
                    row_kind = labels[0][1] if len(labels) == 1 else None
                    kind = tab_kind or row_kind or section or column_tiers.get(column) or "gtd"
                    classification = ("tab explicit" if tab_kind else "row explicit" if row_kind else
                                      "section explicit" if section else
                                      "column explicit" if column in column_tiers else "unlabeled GTD default")
                    evidence.append({"source_id": source["id"], "community": source["title"],
                                     "tab": sheet.title, "row": row_number, "column": column,
                                     "address": address, "kind": kind, "classification": classification})
                for column, cell in enumerate(row, 1):
                    if "0x" in str(cell or "").lower() and len(str(cell).strip()) >= 35 and not wallets(cell):
                        issues.append({"source": source["id"], "tab": sheet.title, "row": row_number,
                                       "column": column, "value": str(cell)[:100], "problem": "invalid wallet"})

    unique = {(row["address"], row["kind"]) for row in evidence}
    manifest = {"reason": "2026-10-19 partner sheet and collab workbook reconciliation",
                "entries": [{"address": address, "kind": kind} for address, kind in sorted(unique)]}
    (args.output_dir / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    with (args.output_dir / "provenance.csv").open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(evidence[0]))
        writer.writeheader()
        writer.writerows(evidence)
    (args.output_dir / "issues.json").write_text(json.dumps(issues, indent=2) + "\n")
    source_counts = {source["id"]: dict(Counter(row["kind"] for row in evidence
                     if row["source_id"] == source["id"])) for source in sources}
    summary = {"sources": len(sources), "source_counts": source_counts,
               "workbook_labeled": dict(workbook_tiers), "evidence_rows": len(evidence),
               "unique_entries": dict(Counter(kind for _, kind in unique)),
               "issues": dict(Counter(row["problem"] for row in issues))}
    (args.output_dir / "summary.json").write_text(json.dumps(summary, indent=2) + "\n")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
